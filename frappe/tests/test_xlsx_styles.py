# Copyright (c) 2025, Frappe Technologies Pvt. Ltd. and Contributors
# License: MIT. See LICENSE

"""
Tests for Excel (XLSX) styling functionality.

This module tests the declarative Excel styling system for reports,
including the XLSXStyleBuilder and make_xlsx with style_config.
"""

import unittest
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils.xlsx_styles import XLSXStyleBuilder
from frappe.utils.xlsxutils import make_xlsx


class TestXLSXStyleBuilder(unittest.TestCase):
	"""Test the XLSXStyleBuilder helper class."""

	def test_register_style_with_dict(self):
		"""Test registering a style with dictionary configuration."""
		builder = XLSXStyleBuilder()
		builder.register_style("header", font={"bold": True, "size": 12, "color": "FFFFFF"})

		self.assertIn("header", builder._named_styles)
		self.assertIsInstance(builder._named_styles["header"]["font"], Font)
		self.assertTrue(builder._named_styles["header"]["font"].bold)
		self.assertEqual(builder._named_styles["header"]["font"].size, 12)
		self.assertEqual(builder._named_styles["header"]["font"].color.rgb, "FFFFFF")

	def test_register_style_with_objects(self):
		"""Test registering a style with openpyxl objects directly."""
		builder = XLSXStyleBuilder()
		font = Font(bold=True, color="FF0000")
		builder.register_style("red_bold", font=font)

		self.assertIn("red_bold", builder._named_styles)
		self.assertIs(builder._named_styles["red_bold"]["font"], font)

	def test_style_column(self):
		"""Test applying style to a column."""
		builder = XLSXStyleBuilder()
		builder.register_style("currency", number_format="#,##0.00")
		builder.style_column(3, "currency")

		config = builder.build()
		self.assertIn(3, config["column_styles"])
		self.assertEqual(config["column_styles"][3]["number_format"], "#,##0.00")

	def test_style_row(self):
		"""Test applying style to a row."""
		builder = XLSXStyleBuilder()
		builder.register_style("total", font={"bold": True})
		builder.style_row(-1, "total")  # Last row

		config = builder.build()
		self.assertIn(-1, config["row_styles"])

	def test_style_cell(self):
		"""Test applying style to a specific cell."""
		builder = XLSXStyleBuilder()
		builder.register_style("highlight", fill={"fgColor": "FFFF00", "fill_type": "solid"})
		builder.style_cell(0, 0, "highlight")

		config = builder.build()
		self.assertIn((0, 0), config["cell_styles"])

	def test_add_conditional_style(self):
		"""Test adding conditional formatting."""
		builder = XLSXStyleBuilder()
		builder.register_style("negative", font={"color": "FF0000"})
		builder.add_conditional_style(lambda r, c, v: isinstance(v, (int, float)) and v < 0, "negative")

		config = builder.build()
		self.assertEqual(len(config["conditional_styles"]), 1)
		self.assertTrue(config["conditional_styles"][0]["condition"](0, 0, -100))
		self.assertFalse(config["conditional_styles"][0]["condition"](0, 0, 100))

	def test_method_chaining(self):
		"""Test that builder methods can be chained."""
		builder = XLSXStyleBuilder()
		result = (
			builder.register_style("s1", font={"bold": True})
			.register_style("s2", font={"italic": True})
			.style_column(0, "s1")
			.style_row(0, "s2")
		)

		self.assertIs(result, builder)
		config = builder.build()
		self.assertIn(0, config["column_styles"])
		self.assertIn(0, config["row_styles"])

	def test_unregistered_style_raises_error(self):
		"""Test that using unregistered style name raises ValueError."""
		builder = XLSXStyleBuilder()

		with self.assertRaises(ValueError) as cm:
			builder.style_column(0, "nonexistent")

		self.assertIn("not registered", str(cm.exception))


class TestMakeXLSXWithStyles(IntegrationTestCase):
	"""Test make_xlsx function with style_config parameter."""

	def test_make_xlsx_without_style_config(self):
		"""Test that make_xlsx works without style_config (backwards compatibility)."""
		data = [["Name", "Amount"], ["Item 1", 100], ["Item 2", 200]]
		xlsx_file = make_xlsx(data, "Test Sheet")

		self.assertIsInstance(xlsx_file, BytesIO)
		wb = load_workbook(xlsx_file)
		self.assertIn("Test Sheet", wb.sheetnames)

	def test_make_xlsx_with_column_style(self):
		"""Test that column styles are applied correctly."""
		data = [["Name", "Amount"], ["Item 1", 100.5], ["Item 2", 200.75]]

		style_config = {
			"column_styles": {1: {"number_format": "#,##0.00"}},
			"row_styles": {},
			"cell_styles": {},
			"conditional_styles": [],
		}

		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Check that the number format is applied to column 1 (Amount)
		# Row 1 is header (0-indexed), row 2 is first data row
		self.assertEqual(ws.cell(2, 2).number_format, "#,##0.00")
		self.assertEqual(ws.cell(3, 2).number_format, "#,##0.00")

	def test_make_xlsx_with_row_style(self):
		"""Test that row styles are applied correctly."""
		data = [["Name", "Amount"], ["Item 1", 100], ["Item 2", 200]]

		style_config = {
			"column_styles": {},
			"row_styles": {0: {"font": Font(bold=True, size=14)}},
			"cell_styles": {},
			"conditional_styles": [],
		}

		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Check that the font is applied to first row
		self.assertTrue(ws.cell(1, 1).font.bold)
		self.assertEqual(ws.cell(1, 1).font.size, 14)

	def test_make_xlsx_with_cell_style(self):
		"""Test that specific cell styles are applied correctly."""
		data = [["Name", "Amount"], ["Item 1", 100], ["Item 2", 200]]

		style_config = {
			"column_styles": {},
			"row_styles": {},
			"cell_styles": {(1, 1): {"fill": PatternFill(fgColor="FFFF00", fill_type="solid")}},
			"conditional_styles": [],
		}

		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Check that the fill is applied to cell (1, 1) which is row 2, col 2 in Excel
		self.assertEqual(ws.cell(2, 2).fill.fgColor.rgb, "FFFF00")

	def test_make_xlsx_with_conditional_style(self):
		"""Test that conditional styles are applied correctly."""
		data = [["Name", "Amount"], ["Item 1", -100], ["Item 2", 200]]

		style_config = {
			"column_styles": {},
			"row_styles": {},
			"cell_styles": {},
			"conditional_styles": [
				{
					"condition": lambda r, c, v: c == 1 and isinstance(v, (int, float)) and v < 0,
					"style": {"font": Font(color="FF0000")},
				}
			],
		}

		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Check that negative value has red font
		self.assertEqual(ws.cell(2, 2).font.color.rgb, "FF0000")
		# Positive value should not have red font
		self.assertNotEqual(ws.cell(3, 2).font.color.rgb, "FF0000")

	def test_make_xlsx_style_precedence(self):
		"""Test that style precedence works correctly (cell > conditional > row > column)."""
		data = [["Name", "Amount"], ["Item 1", 100]]

		# Set different number formats at different levels
		style_config = {
			"column_styles": {1: {"number_format": "0.0"}},
			"row_styles": {1: {"number_format": "0.00"}},
			"cell_styles": {(1, 1): {"number_format": "0.000"}},
			"conditional_styles": [],
		}

		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Cell style should take precedence
		self.assertEqual(ws.cell(2, 2).number_format, "0.000")

	def test_make_xlsx_with_negative_row_index(self):
		"""Test that negative row indices work correctly (e.g., -1 for last row)."""
		data = [["Name", "Amount"], ["Item 1", 100], ["Total", 100]]

		style_config = {
			"column_styles": {},
			"row_styles": {-1: {"font": Font(bold=True)}},  # Last row should be bold
			"cell_styles": {},
			"conditional_styles": [],
		}

		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Last row should be bold
		self.assertTrue(ws.cell(3, 1).font.bold)
		# First data row should not be bold (unless it's a header)
		# Note: First row might be bold from header_index, so check second row
		self.assertFalse(ws.cell(2, 1).font.bold or ws.cell(2, 1).font.bold is None)

	def test_make_xlsx_handles_conditional_style_errors(self):
		"""Test that errors in conditional style functions are handled gracefully."""
		data = [["Name", "Amount"], ["Item 1", 100]]

		# Condition function that will raise an error
		def buggy_condition(r, c, v):
			raise ValueError("Intentional error")

		style_config = {
			"column_styles": {},
			"row_styles": {},
			"cell_styles": {},
			"conditional_styles": [{"condition": buggy_condition, "style": {"font": Font(bold=True)}}],
		}

		# Should not raise an error, just skip the conditional style
		xlsx_file = make_xlsx(data, "Test Sheet", style_config=style_config)
		self.assertIsInstance(xlsx_file, BytesIO)

	def test_make_xlsx_with_builder(self):
		"""Test integration of XLSXStyleBuilder with make_xlsx."""
		data = [["Product", "Q1", "Q2", "Q3", "Total"], ["Widget", 100, 150, 200, 450], ["Gadget", 75, 80, 90, 245]]

		builder = XLSXStyleBuilder()
		builder.register_style("header", font={"bold": True, "size": 12}, fill={"fgColor": "366092", "fill_type": "solid"})
		builder.register_style("currency", number_format="#,##0")
		builder.style_row(0, "header")
		builder.style_column(1, "currency")
		builder.style_column(2, "currency")
		builder.style_column(3, "currency")
		builder.style_column(4, "currency")

		config = builder.build()
		xlsx_file = make_xlsx(data, "Sales Report", style_config=config)

		wb = load_workbook(xlsx_file)
		ws = wb.active

		# Verify header row is bold
		self.assertTrue(ws.cell(1, 1).font.bold)
		# Verify currency format is applied
		self.assertEqual(ws.cell(2, 2).number_format, "#,##0")
		self.assertEqual(ws.cell(2, 5).number_format, "#,##0")
