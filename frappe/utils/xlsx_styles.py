# Copyright (c) 2025, Frappe Technologies Pvt. Ltd. and Contributors
# License: MIT. See LICENSE

"""
Helper classes and utilities for Excel (XLSX) styling in reports.

This module provides a declarative approach to styling Excel exports,
allowing reports to define styling rules without per-cell function calls.
"""

from typing import Any, Callable

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class XLSXStyleBuilder:
	"""
	Helper class to build Excel style configurations declaratively.

	This builder allows reports to define styling rules in a clean, reusable way
	without having to manually construct the style configuration dictionary.

	Example:
		>>> builder = XLSXStyleBuilder()
		>>> builder.register_style("header",
		...     font={"bold": True, "size": 12, "color": "FFFFFF"},
		...     fill={"fgColor": "366092", "fill_type": "solid"})
		>>> builder.style_row(0, "header")
		>>> config = builder.build()
	"""

	def __init__(self):
		"""Initialize the style builder with empty configuration."""
		self.config = {
			"column_styles": {},
			"row_styles": {},
			"cell_styles": {},
			"conditional_styles": [],
		}
		self._named_styles = {}

	def register_style(self, name: str, **kwargs) -> "XLSXStyleBuilder":
		"""
		Register a named style for reuse across multiple cells/rows/columns.

		Args:
			name: Unique name for this style
			font: Font configuration (dict or Font object)
			fill: Fill configuration (dict or PatternFill object)
			number_format: Excel number format string
			alignment: Alignment configuration (dict or Alignment object)
			border: Border configuration (dict or Border object)

		Returns:
			Self for method chaining

		Example:
			>>> builder.register_style("negative",
			...     font={"color": "FF0000", "bold": True},
			...     number_format="#,##0.00")
		"""
		style = {}

		if font_config := kwargs.get("font"):
			if isinstance(font_config, dict):
				style["font"] = Font(**font_config)
			else:
				style["font"] = font_config

		if fill_config := kwargs.get("fill"):
			if isinstance(fill_config, dict):
				style["fill"] = PatternFill(**fill_config)
			else:
				style["fill"] = fill_config

		if "number_format" in kwargs:
			style["number_format"] = kwargs["number_format"]

		if alignment_config := kwargs.get("alignment"):
			if isinstance(alignment_config, dict):
				style["alignment"] = Alignment(**alignment_config)
			else:
				style["alignment"] = alignment_config

		if border_config := kwargs.get("border"):
			if isinstance(border_config, dict):
				style["border"] = Border(**border_config)
			else:
				style["border"] = border_config

		self._named_styles[name] = style
		return self

	def style_column(self, col_idx: int, style_name: str) -> "XLSXStyleBuilder":
		"""
		Apply a named style to an entire column.

		Args:
			col_idx: Zero-based column index
			style_name: Name of previously registered style

		Returns:
			Self for method chaining

		Example:
			>>> builder.style_column(3, "percentage")
		"""
		if style_name not in self._named_styles:
			raise ValueError(f"Style '{style_name}' not registered. Register it first using register_style().")

		self.config["column_styles"][col_idx] = self._named_styles[style_name]
		return self

	def style_row(self, row_idx: int, style_name: str) -> "XLSXStyleBuilder":
		"""
		Apply a named style to an entire row.

		Args:
			row_idx: Zero-based row index (supports negative indices, e.g., -1 for last row)
			style_name: Name of previously registered style

		Returns:
			Self for method chaining

		Example:
			>>> builder.style_row(-1, "total")  # Style last row as total
		"""
		if style_name not in self._named_styles:
			raise ValueError(f"Style '{style_name}' not registered. Register it first using register_style().")

		self.config["row_styles"][row_idx] = self._named_styles[style_name]
		return self

	def style_cell(self, row_idx: int, col_idx: int, style_name: str) -> "XLSXStyleBuilder":
		"""
		Apply a named style to a specific cell.

		Args:
			row_idx: Zero-based row index
			col_idx: Zero-based column index
			style_name: Name of previously registered style

		Returns:
			Self for method chaining

		Example:
			>>> builder.style_cell(0, 0, "header")
		"""
		if style_name not in self._named_styles:
			raise ValueError(f"Style '{style_name}' not registered. Register it first using register_style().")

		self.config["cell_styles"][(row_idx, col_idx)] = self._named_styles[style_name]
		return self

	def add_conditional_style(
		self, condition: Callable[[int, int, Any], bool], style_name: str
	) -> "XLSXStyleBuilder":
		"""
		Add a conditional formatting rule.

		Args:
			condition: Function that takes (row_idx, col_idx, value) and returns True if style should apply
			style_name: Name of previously registered style

		Returns:
			Self for method chaining

		Example:
			>>> builder.add_conditional_style(
			...     lambda r, c, v: c == 5 and isinstance(v, (int, float)) and v < 0,
			...     "negative"
			... )
		"""
		if style_name not in self._named_styles:
			raise ValueError(f"Style '{style_name}' not registered. Register it first using register_style().")

		self.config["conditional_styles"].append(
			{"condition": condition, "style": self._named_styles[style_name]}
		)
		return self

	def build(self) -> dict:
		"""
		Build and return the final style configuration dictionary.

		Returns:
			Dictionary containing column_styles, row_styles, cell_styles, and conditional_styles

		Example:
			>>> config = builder.build()
			>>> # Pass config to make_xlsx(style_config=config)
		"""
		return self.config
